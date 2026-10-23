"""
Bank Reconciliation Automation  (UPI + IMPS + NFS)
==================================================
Reads the ICICI statement + the NPCI settlement files and builds the
calculation table (your "Sheet4") and the voucher table (your "sgeet5").

What is automatic
  * UPI  cycles  (UPI_NTSLMSN<ddmmyy>_<n>C)   -> upi o/w, upi i/w, charges, charges in
  * IMPS cycles  (IMPSNTSLMSN<ddmmyy>_<n>C)   -> imps o/w, imps i/w, charges, charges in
  * NFS  cycles  (NTSLMNL<ddmmyy>_<n>C)       -> atm nfs, atm r, charges, charges in
What is manual (yellow cells in the output - type them in Excel)
  * POS / RuPay lines, DC1 / DC2 lines, and any line the program cannot match.

Everything in the output is a live Excel formula, so the "Diff" column, the
day totals and the voucher update as soon as you fill the yellow cells.

How to run
  1. Put this script in a folder together with the ICICI statement file.
     The settlement files can be in the same folder or in any sub-folders.
  2. python Recon_Automation.py
  3. Open Recon_Calculation.xlsx, fill the yellow cells, check that every
     "Diff" is 0 and every voucher status says OK.
"""

import datetime as dt
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# =====================================================================
# CONFIG  (change only if your files change)
# =====================================================================
STATEMENT_FILE = ""            # "" = auto-detect the statement in this folder
STATEMENT_SHEET_HINT = "icici"  # if the workbook has several sheets, use the one containing this word
OLD_DIFFERENCE = 0.45          # the old difference you keep carrying (editable in the Voucher sheet too)
OUTPUT_NAME = "Recon_Calculation.xlsx"

# Statement column positions (0-based: A=0, B=1 ...). Used ONLY when the
# statement has no header row containing the words Debit / Credit / Balance.
STMT_COLS = {"date": 0, "narration": 1, "debit": 6, "credit": 7, "balance": 8}

# Labels inside the NPCI settlement files (case / spacing insensitive)
HEADER_LABEL = "Description"
SETTLEMENT_CHARGES_LABEL = "Settlement Charges"
FINAL_SETTLEMENT_LABEL = "Final Settlement Amount"
NON_FINANCIAL_MARKER = "FOR NON - FINANCIAL MANDATE/OD TXNs"   # UPI only
AMOUNT_KEYWORD = "Transaction Amount"
TDS_KEYWORD = "TDS"
SUBTOTAL_KEYWORD = "sub total"

TOL = 0.01   # tolerance for money comparisons

# Calc sheet columns (same layout as your Sheet4)
C_UPI_OW, C_UPI_IW, C_ATM, C_ATM_R, C_POS, C_IMPS_OW, C_IMPS_IW, C_CHG, C_CHG_IN = range(3, 12)
C_NET, C_DEBIT, C_CREDIT, C_BAL, C_DIFF, C_SRC, C_NOTE = 13, 14, 15, 16, 17, 18, 19
DEBIT_SIDE = (C_UPI_OW, C_ATM, C_POS, C_IMPS_OW, C_CHG)      # money we pay
CREDIT_SIDE = (C_UPI_IW, C_ATM_R, C_IMPS_IW, C_CHG_IN)       # money we receive

COL_MAP = {                      # kind -> (outward column, inward column)
    "UPI": (C_UPI_OW, C_UPI_IW),
    "IMPS": (C_IMPS_OW, C_IMPS_IW),
    "NFS": (C_ATM, C_ATM_R),
}
MANUAL_COLS = {
    "UPI": [C_UPI_OW, C_UPI_IW, C_CHG, C_CHG_IN],
    "IMPS": [C_IMPS_OW, C_IMPS_IW, C_CHG, C_CHG_IN],
    "NFS": [C_ATM, C_ATM_R, C_CHG, C_CHG_IN],
    "POS": [C_POS],
    "OTHER": list(range(3, 12)),
}

FILE_RE = re.compile(r"^(UPI_NTSLMSN|IMPSNTSLMSN|NTSLMNL)(\d{6})_([A-Za-z0-9]+)\.(xls|xlsx)$", re.I)
KIND_BY_PREFIX = {"UPI_NTSLMSN": "UPI", "IMPSNTSLMSN": "IMPS", "NTSLMNL": "NFS"}
RE_UPI_IMPS = re.compile(r"CMS/(\d{2})\.(\d{2})\.(\d{2})\s+([A-Z0-9]+)\s+(?:DT\s+)?(UPI|IMPS)\b")
RE_NFS = re.compile(r"CMS/NFS\s+([A-Z0-9]+)\s+(\d{6})\b")


def normalize(text) -> str:
    """Lowercase, strip and collapse whitespace so labels match whatever the casing/spacing."""
    return " ".join(str(text).strip().lower().split())


def to_float(v):
    """Number from a cell value; None if blank / not numeric."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if pd.isna(v) else float(v)
    s = str(v).replace(",", "").strip()
    if s == "" or s.lower() in ("nan", "none"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v):
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        return pd.Timestamp(v).date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d-%b-%y", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


# =====================================================================
# 1. SETTLEMENT FILE READERS
# =====================================================================
def _sum(series) -> float:
    return float(pd.to_numeric(series, errors="coerce").fillna(0).sum())


def _find_row(col0, label, after=-1):
    hits = col0.index[(col0 == normalize(label)) & (col0.index > after)]
    return hits[0] if len(hits) else None


def _find_subtotal_row(col0, after):
    hits = col0.index[col0.str.contains(SUBTOTAL_KEYWORD, regex=False) & (col0.index > after)]
    return hits[0] if len(hits) else None


def _block(df, first, last):
    """Rows first..last as a 4-column table (Description / Txns / Debit / Credit)."""
    if last < first:
        out = pd.DataFrame(columns=["Description", "No_of_Txns", "Debit", "Credit"])
        return out
    out = df.loc[first:last].iloc[:, :4].copy()
    out.columns = ["Description", "No_of_Txns", "Debit", "Credit"]
    return out


def read_upi(df):
    """UPI: Amount = 'Transaction Amount' rows above the non-financial marker.
    Fee = Beneficiary/Remitter Sub Totals - Amount (same as the manual method)."""
    col0 = df[0].apply(normalize)
    notes = []
    hdr = _find_row(col0, HEADER_LABEL)
    if hdr is None:
        raise ValueError("header row 'Description' not found")
    sub = _find_subtotal_row(col0, hdr)
    if sub is None:
        raise ValueError("'Sub Totals' row not found")
    marker = _find_row(col0, NON_FINANCIAL_MARKER, hdr)
    if marker is None:
        end = sub - 1
        notes.append("non-financial marker not found, used rows up to Sub Totals")
    else:
        end = marker - 1

    data = _block(df, hdr + 1, end)
    is_amt = data["Description"].apply(lambda x: normalize(AMOUNT_KEYWORD) in normalize(x))
    amt_d = _sum(data.loc[is_amt, "Debit"])
    amt_c = _sum(data.loc[is_amt, "Credit"])
    sub_d = to_float(df.iloc[sub, 2]) or 0.0
    sub_c = to_float(df.iloc[sub, 3]) or 0.0

    # cross-check: every number between header and Sub Totals must add up to the Sub Totals row
    allrows = _block(df, hdr + 1, sub - 1)
    if abs(_sum(allrows["Debit"]) - sub_d) > TOL or abs(_sum(allrows["Credit"]) - sub_c) > TOL:
        notes.append("rows do not add up to Sub Totals - check file")

    return {"deb_amt": amt_d, "cr_amt": amt_c,
            "deb_fee": sub_d - amt_d, "cr_fee": sub_c - amt_c}, notes


def read_imps_style(df):
    """IMPS and NFS: rows between 'Description' and 'Settlement Charges'.
    'Transaction Amount' rows = Amount, every other row = Fee.
    TDS rows found between 'Settlement Charges' and 'Final Settlement Amount' are added to Fee."""
    col0 = df[0].apply(normalize)
    notes = []
    hdr = _find_row(col0, HEADER_LABEL)
    if hdr is None:
        raise ValueError("header row 'Description' not found")
    stl = _find_row(col0, SETTLEMENT_CHARGES_LABEL, hdr)
    if stl is None:
        raise ValueError("'Settlement Charges' row not found")

    # TDS rows (below Settlement Charges, above Final Settlement Amount)
    tds_d = tds_c = 0.0
    fin = _find_row(col0, FINAL_SETTLEMENT_LABEL, stl)
    if fin is not None:
        sec = _block(df, stl + 1, fin - 1)
        is_tds = sec["Description"].apply(lambda x: normalize(TDS_KEYWORD) in normalize(x))
        tds_d = _sum(sec.loc[is_tds, "Debit"])
        tds_c = _sum(sec.loc[is_tds, "Credit"])

    data = _block(df, hdr + 1, stl - 1)
    data = data[data["Description"].apply(lambda x: normalize(x) not in ("", "nan", "none"))]
    is_amt = data["Description"].apply(lambda x: normalize(AMOUNT_KEYWORD) in normalize(x))
    amt_d = _sum(data.loc[is_amt, "Debit"])
    amt_c = _sum(data.loc[is_amt, "Credit"])
    fee_d = _sum(data.loc[~is_amt, "Debit"])
    fee_c = _sum(data.loc[~is_amt, "Credit"])

    # cross-check against the Sub Totals row of the file (TDS is outside the Sub Totals)
    sub = _find_subtotal_row(col0, hdr)
    if sub is not None:
        sub_d = to_float(df.iloc[sub, 2]) or 0.0
        sub_c = to_float(df.iloc[sub, 3]) or 0.0
        if abs(amt_d + fee_d - sub_d) > TOL or abs(amt_c + fee_c - sub_c) > TOL:
            notes.append("rows do not add up to Sub Totals - check file")

    return {"deb_amt": amt_d, "cr_amt": amt_c,
            "deb_fee": fee_d + tds_d, "cr_fee": fee_c + tds_c}, notes


# =====================================================================
# 2. FILE DISCOVERY
# =====================================================================
def index_files(root):
    """Find every settlement file under root (sub-folders included), by file name only."""
    idx = {}
    for dirpath, _dirs, names in os.walk(root):
        for fn in names:
            m = FILE_RE.match(fn)
            if m:
                key = (KIND_BY_PREFIX[m.group(1).upper()], m.group(2), m.group(3).upper())
                idx.setdefault(key, []).append(os.path.join(dirpath, fn))
    return idx


def find_statement(root):
    if STATEMENT_FILE:
        p = STATEMENT_FILE if os.path.isabs(STATEMENT_FILE) else os.path.join(root, STATEMENT_FILE)
        if not os.path.isfile(p):
            sys.exit(f"Statement file not found: {p}")
        return p
    cands = []
    for fn in os.listdir(root):
        low = fn.lower()
        if fn.startswith("~$") or FILE_RE.match(fn) or low.startswith("recon_calculation"):
            continue
        if low.endswith((".xls", ".xlsx", ".csv")):
            cands.append(fn)
    if len(cands) > 1:
        named = [c for c in cands if "icici" in c.lower() or "statement" in c.lower()]
        if len(named) == 1:
            cands = named
    if len(cands) == 1:
        return os.path.join(root, cands[0])
    if not cands:
        sys.exit("No statement file found in this folder. Put the ICICI statement (.xls/.xlsx/.csv) here.")
    sys.exit("More than one possible statement file: " + ", ".join(cands)
             + "\nSet STATEMENT_FILE at the top of this script to the right file name.")


# =====================================================================
# 3. ICICI STATEMENT READER
# =====================================================================
def _detect_columns(raw):
    cols = dict(STMT_COLS)
    for i in range(min(len(raw), 60)):
        cells = [normalize(x) for x in raw.iloc[i].tolist()]
        deb = [j for j, c in enumerate(cells) if c.startswith("debit") or c.startswith("withdrawal")]
        cre = [j for j, c in enumerate(cells) if c.startswith("credit") or c.startswith("deposit")]
        bal = [j for j, c in enumerate(cells) if "balance" in c]
        if deb and cre and bal:
            cols.update(debit=deb[0], credit=cre[0], balance=bal[0])
            nar = [j for j, c in enumerate(cells)
                   if c in ("narration", "description", "particulars", "transaction remarks", "remarks")]
            dte = [j for j, c in enumerate(cells) if "date" in c]
            if nar:
                cols["narration"] = nar[0]
            if dte:
                cols["date"] = dte[0]
            return cols, True
    return cols, False


def read_statement(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        raw = pd.read_csv(path, header=None, dtype=object, engine="python", on_bad_lines="skip")
    else:
        xl = pd.ExcelFile(path)
        sheet = next((s for s in xl.sheet_names if STATEMENT_SHEET_HINT.lower() in s.lower()), xl.sheet_names[0])
        raw = xl.parse(sheet, header=None, dtype=object)

    cols, found = _detect_columns(raw)
    warnings = []
    if not found:
        warnings.append("Statement header row not found - using the default column positions in STMT_COLS.")

    blocks, cur = [], None
    for i in range(len(raw)):
        row = raw.iloc[i].tolist()

        def cell(name):
            j = cols[name]
            return row[j] if j < len(row) else None

        narr = " ".join(str(cell("narration")).split()) if cell("narration") is not None and not pd.isna(cell("narration")) else ""
        date = to_date(cell("date"))
        bal = to_float(cell("balance"))
        deb = to_float(cell("debit")) or 0.0
        cre = to_float(cell("credit")) or 0.0

        if narr.upper().startswith("B/F"):
            cur = {"date": date, "opening": bal, "lines": []}
            blocks.append(cur)
            continue
        if narr == "" or date is None or bal is None:
            continue
        if cur is None:
            cur = {"date": date, "opening": None, "lines": []}
            blocks.append(cur)
        cur["lines"].append({"date": date, "narration": narr, "debit": deb, "credit": cre, "balance": bal})

    blocks = [b for b in blocks if b["lines"]]
    if not blocks:
        sys.exit("No transaction lines could be read from the statement. Check STMT_COLS at the top of the script.")

    for k, b in enumerate(blocks):
        first = b["lines"][0]
        if b["opening"] is None:
            b["opening"] = first["balance"] + first["debit"] - first["credit"]
        b["closing"] = b["lines"][-1]["balance"]
        prev = b["opening"]
        for ln in b["lines"]:
            if abs(prev - ln["debit"] + ln["credit"] - ln["balance"]) > TOL:
                warnings.append(f"Statement balance does not follow on {b['date']:%d-%b-%y}: '{ln['narration'][:40]}' "
                                "(wrong column positions or a missing line?)")
            prev = ln["balance"]
        if k + 1 < len(blocks) and abs(b["closing"] - blocks[k + 1]["opening"]) > TOL:
            warnings.append(f"Closing balance of {b['date']:%d-%b-%y} differs from next day's B/F.")
    return blocks, warnings


# =====================================================================
# 4. MATCH STATEMENT LINES TO FILES
# =====================================================================
def classify(narr):
    n = " ".join(narr.upper().split())
    m = RE_UPI_IMPS.search(n)
    if m:
        return m.group(5), m.group(1) + m.group(2) + m.group(3), m.group(4)
    m = RE_NFS.search(n)
    if m:
        return "NFS", m.group(2), m.group(1)
    if "RUPAY" in n:
        return "POS", None, None
    return "OTHER", None, None


def build_row(line, files, used):
    kind, ddmmyy, cycle = classify(line["narration"])
    row = {"vals": {}, "manual": [], "source": "MANUAL", "notes": [], "kind": kind, "auto": False}
    stmt_net = line["debit"] - line["credit"]

    if kind in COL_MAP and cycle and re.fullmatch(r"\d+C", cycle):
        key = (kind, ddmmyy, cycle)
        paths = files.get(key)
        if not paths:
            row["manual"] = MANUAL_COLS[kind]
            row["notes"].append(f"FILE NOT FOUND for {kind} {ddmmyy} {cycle} - enter manually")
        elif key in used:
            row["manual"] = MANUAL_COLS[kind]
            row["notes"].append("same file already used by an earlier line - enter manually")
        else:
            path = max(paths, key=os.path.getmtime)
            try:
                df = pd.read_excel(path, header=None)
                res, notes = (read_upi if kind == "UPI" else read_imps_style)(df)
                out_c, in_c = COL_MAP[kind]
                vals = {out_c: res["deb_amt"], in_c: res["cr_amt"], C_CHG: res["deb_fee"], C_CHG_IN: res["cr_fee"]}
                row["vals"] = {c: round(v, 2) for c, v in vals.items() if abs(v) > 0.000001}
                row["source"] = os.path.basename(path)
                row["auto"] = True
                used.add(key)
                row["notes"] += notes
                if len(paths) > 1:
                    row["notes"].append("several copies of this file found - newest used")
                if kind == "NFS" and (res["cr_amt"] or res["cr_fee"]):
                    row["notes"].append("NFS file has credit-side values - mapped to 'atm r' / 'charges in', verify")
            except Exception as e:  # keep going, the row just becomes manual
                row["manual"] = MANUAL_COLS[kind]
                row["notes"].append(f"ERROR reading {os.path.basename(path)}: {e}")
    elif kind in COL_MAP:
        row["manual"] = MANUAL_COLS[kind]
        row["notes"].append(f"{kind} {cycle} is entered manually")
    elif kind == "POS":
        row["manual"] = MANUAL_COLS["POS"]
        row["notes"].append("POS / RuPay - enter manually")
    else:
        row["manual"] = MANUAL_COLS["OTHER"]
        row["notes"].append("narration not recognised - enter manually")

    if row["auto"]:
        calc = sum(row["vals"].get(c, 0) for c in DEBIT_SIDE) - sum(row["vals"].get(c, 0) for c in CREDIT_SIDE)
        if abs(calc - stmt_net) > TOL:
            row["notes"].append(f"DOES NOT MATCH STATEMENT: calculated {calc:,.2f} vs statement {stmt_net:,.2f}")
            row["mismatch"] = True
    return row


# =====================================================================
# 5. EXCEL OUTPUT
# =====================================================================
FONT = "Arial"
F_NORM = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=12, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
FILL_HDR = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
FILL_TOT = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
FILL_YEL = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")
FILL_RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
NUM = '#,##0.00;-#,##0.00;-'
HEAD = {3: "upi o/w", 4: "upi i/w", 5: "atm nfs", 6: "atm r", 7: "pos", 8: "imps o/w", 9: "imps i/w",
        10: "charges", 11: "charges in", 13: "Net (calc)", 14: "ICICI Debit", 15: "ICICI Credit",
        16: "ICICI Balance", 17: "Diff", 18: "Source file", 19: "Notes"}


def put(ws, r, c, v, font=F_NORM, fmt=None, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def write_calc(ws, blocks, files):
    used = set()
    r = 1
    put(ws, r, 1, "Calculation table - yellow cells are for manual entry (POS / DC / anything not matched). "
                  "Net = (o/w + atm nfs + pos + charges) - (i/w + atm r + charges in). Diff must be 0.", F_BOLD)
    r = 3
    layout, summary, warnings = [], [], []

    for b in blocks:
        title = f"{b['date']:%d-%b-%y}"
        for c, h in HEAD.items():
            put(ws, r, c, h, F_BOLD, fill=FILL_HDR).alignment = Alignment(horizontal="center")
        put(ws, r, 1, "Date", F_BOLD, fill=FILL_HDR)
        put(ws, r, 2, "Narration", F_BOLD, fill=FILL_HDR)
        r += 1
        bf = r
        put(ws, r, 1, title)
        put(ws, r, 2, "B/F", F_BOLD)
        put(ws, r, C_BAL, round(b["opening"], 2), F_INPUT, NUM)
        r += 1
        first = r
        n_auto = n_manual = n_missing = n_bad = 0

        for ln in b["lines"]:
            row = build_row(ln, files, used)
            put(ws, r, 1, f"{ln['date']:%d-%b-%y}")
            put(ws, r, 2, ln["narration"])
            for c, v in row["vals"].items():
                put(ws, r, c, v, F_INPUT, NUM)
            for c in row["manual"]:
                put(ws, r, c, None, F_INPUT, NUM, FILL_YEL)
            L = get_column_letter
            put(ws, r, C_NET,
                "=(" + "+".join(f"{L(c)}{r}" for c in DEBIT_SIDE) + ")-(" + "+".join(f"{L(c)}{r}" for c in CREDIT_SIDE) + ")",
                F_NORM, NUM)
            put(ws, r, C_DEBIT, ln["debit"], F_INPUT, NUM)
            put(ws, r, C_CREDIT, ln["credit"], F_INPUT, NUM)
            put(ws, r, C_BAL, ln["balance"], F_INPUT, NUM)
            put(ws, r, C_DIFF, f"=ROUND({L(C_NET)}{r}-({L(C_DEBIT)}{r}-{L(C_CREDIT)}{r}),2)", F_NORM, NUM)
            put(ws, r, C_SRC, row["source"])
            put(ws, r, C_NOTE, "; ".join(row["notes"]))
            if row["auto"]:
                n_auto += 1
            else:
                n_manual += 1
                if any("FILE NOT FOUND" in n for n in row["notes"]):
                    n_missing += 1
            if row.get("mismatch"):
                n_bad += 1
            for n in row["notes"]:
                if "Sub Totals" in n or "ERROR" in n or "credit-side" in n or "several copies" in n or "marker" in n:
                    warnings.append(f"{title} | {ln['narration'][:45]} | {n}")
            r += 1
        last = r - 1

        tot = r
        put(ws, r, 1, title, F_BOLD, fill=FILL_TOT)
        put(ws, r, 2, "TOTAL", F_BOLD, fill=FILL_TOT)
        for c in list(range(3, 12)) + [C_NET, C_DEBIT, C_CREDIT]:
            L = get_column_letter(c)
            put(ws, r, c, f"=SUM({L}{first}:{L}{last})", F_BOLD, NUM, FILL_TOT)
        put(ws, r, C_BAL, f"=P{last}", F_BOLD, NUM, FILL_TOT)
        put(ws, r, C_DIFF, f"=SUMPRODUCT(ABS(Q{first}:Q{last}))", F_BOLD, NUM, FILL_TOT)
        layout.append({"title": title, "bf": bf, "first": first, "last": last, "tot": tot})
        summary.append((title, len(b["lines"]), n_auto, n_manual, n_missing, n_bad))
        r += 2

    ws.conditional_formatting.add(f"Q3:Q{r}", CellIsRule(operator="notEqual", formula=["0"], fill=FILL_RED))
    widths = {1: 11, 2: 62, 12: 3, 18: 30, 19: 60}
    for c in range(3, 18):
        widths.setdefault(c, 14)
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "C3"
    return layout, summary, warnings


def write_voucher(ws, layout):
    put(ws, 1, 1, "Voucher table - amounts come from the day TOTAL row of the Calc sheet", F_TITLE)
    put(ws, 2, 1, "Old difference carried forward (edit if it changes)", F_BOLD)
    put(ws, 2, 2, OLD_DIFFERENCE, F_INPUT, '#,##0.00', FILL_YEL)
    r = 4
    for d in layout:
        t, bf, last, tot = d["title"], d["bf"], d["last"], d["tot"]
        put(ws, r, 1, t, F_BOLD, fill=FILL_HDR)
        for c, h in enumerate(["AMOUNT", "GL / CR", "NAME", "BOOKS BALANCE"], start=2):
            put(ws, r, c, h, F_BOLD, fill=FILL_HDR)
        r += 1
        put(ws, r, 1, "Opening = ICICI B/F less old difference")
        put(ws, r, 5, f"=Calc!P{bf}-$B$2", F_LINK, NUM)
        prev = r
        r += 1
        rows = [  # name, amount column in Calc, gl, name2, sign
            ("ATM 8001", C_ATM, 6010, "ICICI", -1),
            ("ATM 8003 pos", C_POS, 6010, "ICICI", -1),
            ("IMPS O/W 8006", C_IMPS_OW, 6010, "ICICI", -1),
            ("UPI O/W 8008", C_UPI_OW, 6010, "ICICI", -1),
            ("ICICI", C_IMPS_IW, 8005, "IMPS I/W", 1),
            ("ICICI", C_UPI_IW, 8007, "UPI I/W", 1),
            ("CHARGES out pl 454", C_CHG, 6010, "ICICI", -1),
            ("CHARGES in gl 6010", C_CHG_IN, "pl 453 in", "", 1),
        ]
        for name, col, gl, name2, sign in rows:
            put(ws, r, 1, name)
            put(ws, r, 2, f"=Calc!{get_column_letter(col)}{tot}", F_LINK, NUM)
            put(ws, r, 3, gl)
            put(ws, r, 4, name2)
            put(ws, r, 5, f"=E{prev}{'+' if sign > 0 else '-'}B{r}", F_NORM, NUM)
            prev = r
            r += 1
        put(ws, r, 1, "Books closing", F_BOLD)
        put(ws, r, 5, f"=E{prev}", F_BOLD, NUM)
        books = r
        r += 1
        put(ws, r, 1, "ICICI closing balance", F_BOLD)
        put(ws, r, 5, f"=Calc!P{last}", F_LINK, NUM)
        icici = r
        r += 1
        put(ws, r, 1, "Difference (ICICI - Books)", F_BOLD)
        put(ws, r, 5, f"=ROUND(E{icici}-E{books},2)", F_BOLD, NUM)
        diff = r
        r += 1
        put(ws, r, 1, "Status", F_BOLD)
        put(ws, r, 5, f'=IF(ROUND(E{diff}-$B$2,2)=0,"OK","CHECK")', F_BOLD)
        ws.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f'E{r}="CHECK"'], fill=FILL_RED))
        r += 1
        put(ws, r, 1, "ATM R total (info - not in voucher, should be 0)")
        put(ws, r, 5, f"=Calc!F{tot}", F_LINK, NUM)
        r += 2
    ws.column_dimensions["A"].width = 48
    for c in "BCDE":
        ws.column_dimensions[c].width = 18


# =====================================================================
# 6. MAIN
# =====================================================================
def main():
    root = os.path.dirname(os.path.abspath(__file__))
    stmt = find_statement(root)
    print(f"Statement : {os.path.basename(stmt)}")
    blocks, warnings = read_statement(stmt)
    files = index_files(root)
    print(f"Settlement files found: {sum(len(v) for v in files.values())}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    layout, summary, w2 = write_calc(ws, blocks, files)
    write_voucher(wb.create_sheet("Voucher"), layout)

    out = os.path.join(root, OUTPUT_NAME)
    try:
        wb.save(out)
    except PermissionError:
        out = os.path.join(root, f"Recon_Calculation_{dt.datetime.now():%H%M%S}.xlsx")
        wb.save(out)

    print("\nDay         | lines | auto | manual | files missing | auto rows not matching statement")
    for t, n, a, m, miss, bad in summary:
        print(f"{t:<11} | {n:>5} | {a:>4} | {m:>6} | {miss:>13} | {bad:>5}")
    for w in warnings + w2:
        print("WARNING:", w)
    print(f"\nSaved: {out}")
    print("Next: fill the yellow cells in Excel, then check Diff = 0 and Voucher status = OK.")


if __name__ == "__main__":
    main()
